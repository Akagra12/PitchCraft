import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(40, 44, 52)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    i = 0
    in_code_block = False
    code_block_lines = []

    while i < len(lines):
        line = lines[i].rstrip('\n')

        # Code block
        if line.startswith('```'):
            if in_code_block:
                # End code block
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, 'F4F6F9')
                set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run('\n'.join(code_block_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(33, 37, 41)
                doc.add_paragraph() # Spacing
                code_block_lines = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Markdown Table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Parse table rows
            rows_data = []
            for tl in table_lines:
                # Skip separator lines like |---|---|
                cleaned = tl.strip('|')
                parts = [p.strip() for p in cleaned.split('|')]
                if all(re.match(r'^:?-+:?$', p) for p in parts if p):
                    continue
                rows_data.append(parts)

            if rows_data:
                num_cols = max(len(r) for r in rows_data)
                table = doc.add_table(rows=len(rows_data), cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                for r_idx, row in enumerate(rows_data):
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell = table.cell(r_idx, c_idx)
                        text = row[c_idx] if c_idx < len(row) else ''
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        
                        if is_header:
                            set_cell_background(cell, '2D3748')
                        elif r_idx % 2 == 1:
                            set_cell_background(cell, 'F7FAFC')
                        else:
                            set_cell_background(cell, 'FFFFFF')

                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        run = p.add_run(text)
                        if is_header:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.font.size = Pt(10)
                        else:
                            run.font.size = Pt(9.5)
                
                doc.add_paragraph() # Spacing
            continue

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Horizontal Rule
        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            h = doc.add_heading(line[2:].strip(), level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(6)
            for r in h.runs:
                r.font.name = 'Arial'
                r.font.color.rgb = RGBColor(26, 32, 44)
                r.font.size = Pt(20)
                r.font.bold = True
            i += 1
            continue

        if line.startswith('## '):
            h = doc.add_heading(line[3:].strip(), level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = 'Arial'
                r.font.color.rgb = RGBColor(43, 108, 176)
                r.font.size = Pt(15)
                r.font.bold = True
            i += 1
            continue

        if line.startswith('### '):
            h = doc.add_heading(line[4:].strip(), level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(2)
            for r in h.runs:
                r.font.name = 'Arial'
                r.font.color.rgb = RGBColor(45, 55, 72)
                r.font.size = Pt(12)
                r.font.bold = True
            i += 1
            continue

        # Blockquote
        if line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line[2:].strip())
            r.font.italic = True
            r.font.color.rgb = RGBColor(113, 128, 150)
            i += 1
            continue

        # Unordered List
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            item_text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            # Simple bold formatting parser
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Numbered List
        m_num = re.match(r'^(\d+)\.\s+(.*)$', line.strip())
        if m_num:
            item_text = m_num.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.font.bold = True
            else:
                p.add_run(part)
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated DOCX at: {docx_path}")

if __name__ == '__main__':
    src = r"c:\Users\AKAGRA\Desktop\project\The Startup Idea Evaluator(Mobile App Internship Assignment)\PitchCraft\PROJECT_DOCUMENT.md"
    dst = r"c:\Users\AKAGRA\Desktop\project\The Startup Idea Evaluator(Mobile App Internship Assignment)\PitchCraft\PitchCraft_Project_Document.docx"
    markdown_to_docx(src, dst)
